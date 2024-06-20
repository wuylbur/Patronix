#this code is part of the Final Degree Project in UNIR by Gonzalo Castro

#This library is to read/write Excel  xlsx/xlsm/xltx/xltm files.
import openpyxl

#This library is for creating and updating Microsoft Word (.docx) files
import docx

#This library is for creating, reading, and updating PowerPoint (.pptx) files.
from pptx import Presentation

#This is a library to read and write OpenDocument v. 1.2 files
from odf.opendocument import load
from odf.meta import UserDefined


def add_Metadata_GDPR (p1,e1,result_GDPR):
  """
  add the results found to the document metadata,
  differentiating by document type

  """
  OK_Meta= False  
  
  for result in result_GDPR:
    comment = result[0]
          
    if e1 in ("DOCX", "DOC"):  #if the document is  WORD
      try:   
        file_meta= docx.Document(p1)
        properties = file_meta.core_properties
        
        if properties.comments is not None:
          properties.comments += "\n" + comment
        else:
          properties.comments = comment
        
        file_meta.save(p1)
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")
        
    elif e1 in  ("XLS","XLSX"): #if the document is Excel 
      try:
        book = openpyxl.load_workbook(p1)
        properties = book.properties
        
        if properties.description is not None:
          properties.description += "\n" + comment
        else:
          properties.description = comment
        
        book.save(p1)
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")
        
    elif e1 in ("PPT", "PPTX"): #if the document is Power Point
      try:
        presentation_tmp = Presentation(p1)  
        properties = presentation_tmp.core_properties
        properties.comments = comment
        presentation_tmp.save(p1)
        
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")

    elif e1 in ("ODS", "ODT", "ODP"):  #if the document is Open Office
      try:   
        doc = load(p1)
        comment2= UserDefined(name="Comment", text=comment)
        doc.meta.addElement(comment2)
        doc.save(p1)
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")   
   
  return OK_Meta  
  
def add_Metadata_Pattern (p1,e1,pattern_value):
  #add the results found to the document metadata,
  #differentiating by document type
  
  
  OK_Meta=False
  
  if e1 in ("DOCX", "DOC"):  #if the document is  WORD
      try:   
        file_meta= docx.Document(p1)
        properties = file_meta.core_properties
        
        if properties.comments is not None:
          properties.comments += "\n" + pattern_value
        else:
          properties.comments = pattern_value
        
        file_meta.save(p1)
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")
        
  elif e1 in  ("XLS","XLSX"):#if the document is  EXCEL
    try:
      book = openpyxl.load_workbook(p1)
      properties = book.properties
      
      if properties.description is not None:
        properties.description += "\n" + pattern_value
      else:
        properties.description = pattern_value
      
      book.save(p1)
      OK_Meta=True
      
      
    except Exception as e:
      print(f"Error: {e}")
      
      
  elif e1 in ("PPT", "PPTX"): #if the document is  POWER POINT
    try:
      presentation_tmp = Presentation(p1)  
      properties = presentation_tmp.core_properties
      properties.comments = pattern_value
      presentation_tmp.save(p1)
      
      OK_Meta=True
      
      
    except Exception as e:
      print(f"Error: {e}")
    
  elif e1 in ("ODS", "ODT", "ODP"):  #if the document is  Open Office
      try:   
        doc = load(p1)
        comment2= UserDefined(name="Comment", text=pattern_value)
        doc.meta.addElement(comment2)
        doc.save(p1)
        OK_Meta=True
        
        
      except Exception as e:
        print(f"Error: {e}")   
   
  return OK_Meta

if __name__ == "__main__":
    add_Metadata_GDPR()
    add_Metadata_Pattern()